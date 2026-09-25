import shutil
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ML = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ML))

from preprocess import FILA_ENCABEZADO, FILA_INICIO, HOJA_DATOS

PLANTILLAS = ML.parent / "fichas"
ARCHIVOS = {
    1: "ficha_1_rendimiento.xlsx",
    2: "ficha_2_asistencia.xlsx",
    3: "ficha_3_reuniones.xlsx",
}


def llenar_ficha(destino: Path, ficha: int, filas: list[dict], **encabezado) -> Path:
    """Copia la plantilla vacía y escribe filas de prueba por nombre de columna.

    encabezado admite anio (B2) y escala (B7, solo ficha 1).
    """
    ruta = destino / ARCHIVOS[ficha]
    shutil.copy(PLANTILLAS / ARCHIVOS[ficha], ruta)
    libro = load_workbook(ruta)
    hoja = libro[HOJA_DATOS]
    if "anio" in encabezado:
        hoja["B2"] = encabezado["anio"]
    if "escala" in encabezado:
        hoja["B7"] = encabezado["escala"]
    columnas = {
        celda.value: celda.column for celda in hoja[FILA_ENCABEZADO] if celda.value
    }
    for i, fila in enumerate(filas):
        for nombre, valor in fila.items():
            hoja.cell(FILA_INICIO + i, columnas[nombre], valor)
    libro.save(ruta)
    return ruta


def alumno(codigo="EST-001", grado=4, momento="pre", anio=2026, **extra) -> dict:
    return {
        "codigo": codigo,
        "grado": grado,
        "seccion": "A",
        "momento": momento,
        "anio": anio,
        **extra,
    }


NOTAS_142 = dict(
    zip(
        [
            "matematica",
            "comunicacion",
            "ingles",
            "arte_cultura",
            "ciencias_sociales",
            "dpcc",
            "educacion_fisica",
            "educacion_religiosa",
            "ciencia_tecnologia",
            "ept",
        ],
        [14, 15, 13, 16, 12, 14, 15, 13, 14, 16],
    )
)


@pytest.fixture
def lote(tmp_path):
    """Crea un lote con las 3 fichas; cada argumento es la lista de filas."""

    def crear(f1, f2, f3, **encabezado):
        encabezado_f1 = dict(encabezado)
        encabezado_otras = {k: v for k, v in encabezado.items() if k != "escala"}
        llenar_ficha(tmp_path, 1, f1, **encabezado_f1)
        llenar_ficha(tmp_path, 2, f2, **encabezado_otras)
        llenar_ficha(tmp_path, 3, f3, **encabezado_otras)
        return tmp_path

    return crear
