"""
generar_fichas.py

Genera las 3 fichas de registro de datos (instrumento de la tesis) como
plantillas Excel vacías, una por dimensión de la variable dependiente:

  - ficha_1_rendimiento.xlsx -> Promedio = Σ Notas / n        [1]
  - ficha_2_asistencia.xlsx  -> Asistencia = (DA / DP) × 100  [1]
  - ficha_3_reuniones.xlsx   -> CA = (RA / RT) × 100          [39]

Los nombres de columna coinciden con el CSV de carga masiva de
docs/VARIABLES.md. Las plantillas no contienen datos: una ficha llena es un
dato real y se guarda en data/raw/ (gitignored), nunca en esta carpeta.

Uso (desde la raíz del repo, con el entorno de ml/):
    python fichas/generar_fichas.py
    python fichas/generar_fichas.py --filas 200 --out fichas
"""

import argparse
import os
from collections.abc import Callable
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

# Distribución de la hoja "Datos"
FILA_TITULO = 1
FILA_ESCALA = 7  # solo ficha 1
FILA_ENCABEZADO = 9
FILA_INICIO = 10
CELDA_ESCALA = f"$B${FILA_ESCALA}"

# Decisión D3 (docs/VARIABLES.md)
PERIODO_CORTE = "I bimestre"

AZUL = "1F4E78"
GRIS = "E7E6E6"
AMBAR = "FFE699"
ROJO = "F4B6B6"
BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

# Áreas curriculares de secundaria (Currículo Nacional, MINEDU)
AREAS = [
    ("matematica", "Matemática"),
    ("comunicacion", "Comunicación"),
    ("ingles", "Inglés como lengua extranjera"),
    ("arte_cultura", "Arte y Cultura"),
    ("ciencias_sociales", "Ciencias Sociales"),
    ("dpcc", "Desarrollo Personal, Ciudadanía y Cívica"),
    ("educacion_fisica", "Educación Física"),
    ("educacion_religiosa", "Educación Religiosa"),
    ("ciencia_tecnologia", "Ciencia y Tecnología"),
    ("ept", "Educación para el Trabajo"),
]

# Conversión provisional de escala literal (decisión D1, docs/VARIABLES.md)
CONVERSION_LITERAL = [
    ("AD", 4, "Logro destacado"),
    ("A", 3, "Logro esperado"),
    ("B", 2, "En proceso"),
    ("C", 1, "En inicio"),
]
HOJA_CONVERSION = "Conversión"
RANGO_CODIGOS = f"'{HOJA_CONVERSION}'!$A$5:$A$8"
RANGO_VALORES = f"'{HOJA_CONVERSION}'!$B$5:$B$8"

MOTIVOS_EXCLUSION = "traslado definitivo,dimensión incompleta,registro inconsistente"


@dataclass
class Columna:
    nombre: str
    descripcion: str
    ancho: int = 12
    # Si tiene fórmula es calculada (bloqueada, gris); si no, es de captura.
    formula: Callable[[int, dict], str] | None = None
    # Recibe las referencias de la primera fila: {"c": celda, "<columna>": celda}
    validacion: Callable[[dict], DataValidation] | None = None
    regla: str = ""
    obligatoria: bool = False


@dataclass
class Ficha:
    archivo: str
    titulo: str
    dimension: str
    indicador: str
    formula_tesis: str
    cita: str
    columnas: list[Columna]
    notas: list[str] = field(default_factory=list)
    # Formato condicional extra: (columna, plantilla de fórmula, color)
    alertas: list[tuple[str, str, str]] = field(default_factory=list)
    con_escala: bool = False


# --- Validaciones -----------------------------------------------------------


def _dv(tipo: str, error: str, **kwargs) -> DataValidation:
    dv = DataValidation(type=tipo, allow_blank=True, **kwargs)
    dv.showErrorMessage = True
    dv.errorTitle = "Valor no válido"
    dv.error = error
    return dv


def dv_lista(valores: str, error: str) -> Callable[[dict], DataValidation]:
    return lambda _: _dv("list", error, formula1=f'"{valores}"')


def dv_entero(minimo: int, maximo: int, error: str) -> Callable[[dict], DataValidation]:
    return lambda _: _dv(
        "whole", error, operator="between", formula1=str(minimo), formula2=str(maximo)
    )


def dv_formula(plantilla: str, error: str) -> Callable[[dict], DataValidation]:
    """plantilla usa {c} para la propia celda y {<columna>} para otras columnas."""
    return lambda refs: _dv("custom", error, formula1=plantilla.format(**refs))


# --- Columnas comunes -------------------------------------------------------


def columnas_identificacion() -> list[Columna]:
    return [
        Columna(
            "codigo",
            "Código anonimizado del estudiante. Sin nombres ni DNI.",
            ancho=11,
            validacion=dv_formula(
                'AND(LEN({c})=7,LEFT({c},4)="EST-",ISNUMBER(VALUE(RIGHT({c},3))))',
                "Use el formato EST-001 (sin nombres ni DNI).",
            ),
            regla="Formato EST-### (ej. EST-001)",
            obligatoria=True,
        ),
        Columna(
            "grado",
            "Grado de secundaria.",
            ancho=8,
            validacion=dv_entero(3, 4, "El grado debe ser 3 o 4."),
            regla="3 o 4",
            obligatoria=True,
        ),
        Columna(
            "seccion",
            "Sección (ej. A).",
            ancho=9,
            validacion=lambda _: _dv(
                "textLength",
                "La sección tiene 1 o 2 caracteres (ej. A).",
                operator="between",
                formula1="1",
                formula2="2",
            ),
            regla="1–2 caracteres",
            obligatoria=True,
        ),
        Columna(
            "grupo",
            "Grupo del diseño cuasiexperimental: 3.° = control, 4.° = experimental.",
            ancho=13,
            formula=lambda r, c: (
                f'=IF({c["grado"]}{r}=3,"control",'
                f'IF({c["grado"]}{r}=4,"experimental",""))'
            ),
            regla="Se calcula a partir del grado",
        ),
        Columna(
            "momento",
            "Momento de la medición.",
            ancho=10,
            validacion=dv_lista("pre,post", "Elija pre o post."),
            regla="pre o post",
            obligatoria=True,
        ),
    ]


def columnas_cierre() -> list[Columna]:
    return [
        Columna(
            "anio",
            "Año lectivo del registro. 2024–2025 = histórico de entrenamiento.",
            ancho=8,
            validacion=dv_entero(2024, 2026, "El año debe estar entre 2024 y 2026."),
            regla="2024, 2025 o 2026",
            obligatoria=True,
        ),
        Columna(
            "deserto",
            "Solo histórico 2024–2025: 1 si abandonó ese año, 0 si no. "
            "Dejar vacío en 2026.",
            ancho=9,
            validacion=dv_entero(0, 1, "Use 1 (desertó) o 0 (no desertó)."),
            regla="0 o 1; vacío en 2026",
        ),
        Columna(
            "excluido",
            "Marcar sí si cumple un criterio de exclusión de la tesis.",
            ancho=9,
            validacion=dv_lista("sí,no", "Elija sí o no."),
            regla="sí o no",
        ),
        Columna(
            "motivo_exclusion",
            "Obligatorio si excluido = sí.",
            ancho=22,
            validacion=dv_lista(MOTIVOS_EXCLUSION, "Elija un motivo de la lista."),
            regla=MOTIVOS_EXCLUSION.replace(",", " / "),
        ),
        Columna("observaciones", "Texto libre (sin datos personales).", ancho=30),
    ]


# --- Definición de las 3 fichas ---------------------------------------------


def ficha_rendimiento() -> Ficha:
    notas = [
        Columna(
            nombre,
            f"Calificación del periodo en {etiqueta}. Vacío si no aplica "
            "(ej. exonerado).",
            ancho=11,
            validacion=dv_formula(
                f'IF({CELDA_ESCALA}="literal",'
                'OR({c}="AD",{c}="A",{c}="B",{c}="C"),'
                "AND(ISNUMBER({c}),{c}>=0,{c}<=20))",
                "Vigesimal: número de 0 a 20. Literal: AD, A, B o C.",
            ),
            regla="0–20 (vigesimal) o AD/A/B/C (literal)",
        )
        for nombre, etiqueta in AREAS
    ]
    primera, ultima = AREAS[0][0], AREAS[-1][0]

    def rango(r: int, c: dict) -> str:
        return f"{c[primera]}{r}:{c[ultima]}{r}"

    def conteo(r: int, c: dict) -> str:
        return (
            f'IF({CELDA_ESCALA}="literal",'
            f"SUMPRODUCT(COUNTIF({rango(r, c)},{RANGO_CODIGOS})),"
            f'COUNTIFS({rango(r, c)},">=0",{rango(r, c)},"<=20"))'
        )

    calculadas = [
        Columna(
            "n_notas",
            "n: cantidad de calificaciones válidas registradas.",
            ancho=9,
            formula=lambda r, c: f'=IF({conteo(r, c)}=0,"",{conteo(r, c)})',
            regla="Cuenta solo notas válidas",
        ),
        Columna(
            "suma_notas",
            "Σ Notas. En escala literal suma los valores de la hoja Conversión.",
            ancho=11,
            formula=lambda r, c: (
                f'=IF({c["n_notas"]}{r}="","",IF({CELDA_ESCALA}="literal",'
                f"SUMPRODUCT(COUNTIF({rango(r, c)},{RANGO_CODIGOS})*{RANGO_VALORES}),"
                f'SUMIFS({rango(r, c)},{rango(r, c)},">=0",{rango(r, c)},"<=20")))'
            ),
            regla="Σ Notas",
        ),
        Columna(
            "promedio",
            "Promedio = Σ Notas / n, redondeado a 2 decimales.",
            ancho=10,
            formula=lambda r, c: (
                f'=IF({c["n_notas"]}{r}="","",'
                f"ROUND({c['suma_notas']}{r}/{c['n_notas']}{r},2))"
            ),
            regla="Σ Notas / n",
        ),
    ]
    alertas = [
        (
            nombre,
            (
                f'AND({{c}}<>"",NOT(IF({CELDA_ESCALA}="literal",'
                'OR({c}="AD",{c}="A",{c}="B",{c}="C"),'
                "AND(ISNUMBER({c}),{c}>=0,{c}<=20))))"
            ),
            ROJO,
        )
        for nombre, _ in AREAS
    ]
    return Ficha(
        archivo="ficha_1_rendimiento.xlsx",
        titulo="Ficha 1 — Rendimiento académico",
        dimension="Rendimiento académico",
        indicador="Promedio de calificaciones",
        formula_tesis="Promedio = Σ Notas / n",
        cita="Albonny y Duru [1]",
        columnas=columnas_identificacion() + notas + calculadas + columnas_cierre(),
        notas=[
            (
                "Registre la calificación de cada área al cierre del periodo de corte. "
                "Deje vacía el área que no aplique (ej. exoneración); n cuenta solo "
                "las notas registradas."
            ),
            (
                "Elija la escala en la celda B7. Con escala literal, AD/A/B/C se "
                "convierten según la hoja Conversión (provisional, decisión D1)."
            ),
            (
                "Una nota no válida para la escala elegida se marca en rojo y no se "
                "cuenta en el promedio."
            ),
        ],
        alertas=alertas,
        con_escala=True,
    )


def ficha_porcentaje(
    archivo: str,
    titulo: str,
    dimension: str,
    indicador: str,
    formula_tesis: str,
    cita: str,
    programado: tuple[str, str, str],
    realizado: tuple[str, str, str],
    calculado: tuple[str, str],
    notas: list[str],
) -> Ficha:
    """Ficha de la forma indicador = (realizado / programado) × 100."""
    n_prog, sim_prog, desc_prog = programado
    n_real, sim_real, desc_real = realizado
    n_calc, desc_calc = calculado
    columnas = [
        Columna(
            n_prog,
            f"{sim_prog}: {desc_prog}",
            ancho=18,
            validacion=dv_entero(1, 400, f"{sim_prog} debe ser un entero mayor que 0."),
            regla=f"{sim_prog}: entero > 0",
            obligatoria=True,
        ),
        Columna(
            n_real,
            f"{sim_real}: {desc_real}",
            ancho=18,
            validacion=dv_formula(
                "AND(ISNUMBER({c}),{c}=INT({c}),{c}>=0,{c}<={" + n_prog + "})",
                f"{sim_real} debe ser un entero entre 0 y {sim_prog}.",
            ),
            regla=f"{sim_real}: entero, 0 ≤ {sim_real} ≤ {sim_prog}",
            obligatoria=True,
        ),
        Columna(
            n_calc,
            desc_calc,
            ancho=15,
            formula=lambda r, c: (
                f'=IF(OR({c[n_prog]}{r}="",{c[n_real]}{r}="",{c[n_prog]}{r}<=0),"",'
                f"ROUND({c[n_real]}{r}/{c[n_prog]}{r}*100,2))"
            ),
            regla=formula_tesis,
        ),
    ]
    return Ficha(
        archivo=archivo,
        titulo=titulo,
        dimension=dimension,
        indicador=indicador,
        formula_tesis=formula_tesis,
        cita=cita,
        columnas=columnas_identificacion() + columnas + columnas_cierre(),
        notas=notas,
        alertas=[
            (n_real, f'AND({{c}}<>"",{{c}}>{{{n_prog}}})', ROJO),
        ],
    )


def fichas() -> list[Ficha]:
    return [
        ficha_rendimiento(),
        ficha_porcentaje(
            archivo="ficha_2_asistencia.xlsx",
            titulo="Ficha 2 — Asistencia escolar",
            dimension="Asistencia escolar",
            indicador="Porcentaje de asistencia",
            formula_tesis="Asistencia = (DA / DP) × 100",
            cita="Albonny y Duru [1]",
            programado=(
                "dias_programados",
                "DP",
                "días lectivos programados en el periodo de corte.",
            ),
            realizado=(
                "dias_asistidos",
                "DA",
                "días que el estudiante asistió en el periodo de corte.",
            ),
            calculado=("pct_asistencia", "Asistencia = (DA / DP) × 100, 2 decimales."),
            notas=[
                (
                    "Cuente solo los días lectivos del periodo de corte (no el año "
                    "completo). Las tardanzas cuentan como asistencia; las faltas "
                    "justificadas y no justificadas, como inasistencia."
                ),
            ],
        ),
        ficha_porcentaje(
            archivo="ficha_3_reuniones.xlsx",
            titulo="Ficha 3 — Apoyo familiar",
            dimension="Apoyo familiar",
            indicador="Asistencia a reuniones de padres",
            formula_tesis="CA = (RA / RT) × 100",
            cita="Ttito y Choque [39]",
            programado=(
                "reuniones_programadas",
                "RT",
                "reuniones de padres convocadas en el periodo de corte.",
            ),
            realizado=(
                "reuniones_asistidas",
                "RA",
                "reuniones a las que asistió el padre, madre o apoderado.",
            ),
            calculado=("pct_reuniones", "CA = (RA / RT) × 100, 2 decimales."),
            notas=[
                (
                    "Cuente las reuniones convocadas por la IE o el tutor para el "
                    "aula del estudiante dentro del periodo de corte, según el "
                    "registro de asistencia de padres."
                ),
            ],
        ),
    ]


# --- Construcción del libro -------------------------------------------------

ESTILO_TITULO = Font(bold=True, size=14, color=AZUL)
ESTILO_ENCABEZADO = Font(bold=True, color="FFFFFF")
RELLENO_ENCABEZADO = PatternFill("solid", fgColor=AZUL)
RELLENO_CALCULADA = PatternFill("solid", fgColor=GRIS)
LIBRE = Protection(locked=False)


def relleno_alerta(color: str) -> PatternFill:
    """En formato condicional Excel pinta con bgColor, no con fgColor."""
    return PatternFill("solid", fgColor=color, bgColor=color)


def escribir_metadatos(ws: Worksheet, ficha: Ficha) -> None:
    ws.cell(FILA_TITULO, 1, ficha.titulo).font = ESTILO_TITULO
    ws.cell(FILA_TITULO, 5, f"{ficha.formula_tesis}  ·  {ficha.cita}").font = Font(
        italic=True, color=AZUL
    )
    campos = [
        ("Año lectivo", dv_entero(2024, 2026, "Año entre 2024 y 2026.")),
        ("Periodo de corte", dv_lista(PERIODO_CORTE, "El corte es el I bimestre.")),
        ("Fecha de inicio del periodo", None),
        ("Fecha de fin del periodo", None),
        ("Responsable del llenado (cargo)", None),
    ]
    if ficha.con_escala:
        campos.append(
            ("Escala de calificación", dv_lista("vigesimal,literal", "Elija una."))
        )
    for i, (etiqueta, validacion) in enumerate(campos, start=2):
        ws.cell(i, 1, etiqueta).font = Font(bold=True)
        valor = ws.cell(i, 2)
        valor.protection = LIBRE
        valor.border = BORDE
        valor.fill = PatternFill("solid", fgColor="FFFFFF")
        if "Fecha" in etiqueta:
            valor.number_format = "DD/MM/YYYY"
        if validacion:
            dv = validacion({"c": valor.coordinate})
            ws.add_data_validation(dv)
            dv.add(valor.coordinate)
    ws["B3"] = PERIODO_CORTE
    if ficha.con_escala:
        ws[f"B{FILA_ESCALA}"] = "vigesimal"
    ws.column_dimensions["A"].width = 30


def escribir_tabla(ws: Worksheet, ficha: Ficha, filas: int) -> None:
    letras = {
        col.nombre: get_column_letter(i) for i, col in enumerate(ficha.columnas, 1)
    }
    refs = {nombre: f"{letra}{FILA_INICIO}" for nombre, letra in letras.items()}
    ultima_fila = FILA_INICIO + filas - 1
    codigo = f"${letras['codigo']}"

    for i, col in enumerate(ficha.columnas, start=1):
        letra = letras[col.nombre]
        cab = ws.cell(FILA_ENCABEZADO, i, col.nombre)
        cab.font = ESTILO_ENCABEZADO
        cab.fill = (
            RELLENO_ENCABEZADO
            if not col.formula
            else PatternFill("solid", fgColor="595959")
        )
        cab.alignment = Alignment(horizontal="center", vertical="center")
        ancho = max(col.ancho, len(col.nombre) + 2)
        if i > 1:  # la columna A también lleva las etiquetas de metadatos
            ws.column_dimensions[letra].width = ancho

        for r in range(FILA_INICIO, ultima_fila + 1):
            celda = ws.cell(r, i)
            celda.border = BORDE
            if col.formula:
                celda.value = col.formula(r, letras)
                celda.fill = RELLENO_CALCULADA
            else:
                celda.protection = LIBRE

        rango = f"{letra}{FILA_INICIO}:{letra}{ultima_fila}"
        primera = f"{letra}{FILA_INICIO}"
        if col.validacion:
            dv = col.validacion({**refs, "c": primera})
            ws.add_data_validation(dv)
            dv.add(rango)
        if col.obligatoria:
            # Ámbar: falta un dato obligatorio en una fila con código
            ws.conditional_formatting.add(
                rango,
                FormulaRule(
                    formula=[f'AND({codigo}{FILA_INICIO}<>"",{primera}="")'],
                    fill=relleno_alerta(AMBAR),
                ),
            )

    # Ámbar: excluido = sí sin motivo
    motivo = f"{letras['motivo_exclusion']}{FILA_INICIO}"
    ws.conditional_formatting.add(
        f"{letras['motivo_exclusion']}{FILA_INICIO}:"
        f"{letras['motivo_exclusion']}{ultima_fila}",
        FormulaRule(
            formula=[f'AND(${letras["excluido"]}{FILA_INICIO}="sí",{motivo}="")'],
            fill=relleno_alerta(AMBAR),
        ),
    )
    # Rojo: valores inconsistentes propios de cada ficha
    for nombre, plantilla, color in ficha.alertas:
        letra = letras[nombre]
        formula = plantilla.format(**{**refs, "c": refs[nombre]})
        ws.conditional_formatting.add(
            f"{letra}{FILA_INICIO}:{letra}{ultima_fila}",
            FormulaRule(formula=[formula], fill=relleno_alerta(color)),
        )

    ultima_col = get_column_letter(len(ficha.columnas))
    ws.auto_filter.ref = f"A{FILA_ENCABEZADO}:{ultima_col}{ultima_fila}"
    ws.freeze_panes = f"B{FILA_INICIO}"


def proteger(ws: Worksheet) -> None:
    """Bloquea las fórmulas sin contraseña: se puede desproteger desde Revisar."""
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.autoFilter = False
    ws.protection.sort = False


def hoja_instrucciones(wb: Workbook, ficha: Ficha) -> None:
    ws = wb.create_sheet("Instrucciones", 0)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 62
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 36
    envolver = Alignment(wrap_text=True, vertical="top")

    ws["A1"] = ficha.titulo
    ws["A1"].font = ESTILO_TITULO
    ficha_tesis = [
        ("Dimensión", ficha.dimension),
        ("Indicador", ficha.indicador),
        ("Fórmula (tesis)", f"{ficha.formula_tesis}  ·  {ficha.cita}"),
        (
            "Periodo de corte",
            (
                "Datos acumulados hasta el cierre del I bimestre. El "
                "mismo corte se usa para el PRE 2026 y para el histórico 2024–2025, "
                "para que el modelo no aprenda de la ausencia de quien ya desertó."
            ),
        ),
        (
            "Privacidad",
            (
                "Solo códigos EST-###; ningún nombre ni DNI (Ley N.° 29733). La tabla "
                "código ↔ nombre la conserva la IE, fuera del sistema. La ficha llena "
                "se guarda en data/raw/ y no se sube al repositorio."
            ),
        ),
        (
            "Exclusión",
            (
                "Si falta alguna de las 3 dimensiones o hubo traslado definitivo antes "
                "del cierre, marque excluido = sí y el motivo. No borre la fila: los "
                "excluidos se reportan en Resultados."
            ),
        ),
        (
            "Colores",
            (
                "Gris: calculada, no editar. Ámbar: falta un dato obligatorio. Rojo: "
                "valor inconsistente."
            ),
        ),
    ]
    fila = 3
    for etiqueta, texto in ficha_tesis + [("Nota", n) for n in ficha.notas]:
        ws.cell(fila, 1, etiqueta).font = Font(bold=True)
        ws.cell(fila, 2, texto).alignment = envolver
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
        ws.row_dimensions[fila].height = 15 * max(1, len(texto) // 95 + 1)
        fila += 1

    fila += 1
    ws.cell(fila, 1, "Diccionario de columnas (hoja Datos)").font = Font(
        bold=True, color=AZUL
    )
    fila += 1
    for i, cab in enumerate(["Columna", "Descripción", "Tipo", "Regla"], start=1):
        celda = ws.cell(fila, i, cab)
        celda.font = ESTILO_ENCABEZADO
        celda.fill = RELLENO_ENCABEZADO
    for col in ficha.columnas:
        fila += 1
        tipo = "calculada" if col.formula else "captura"
        if col.obligatoria:
            tipo += " *"
        for i, valor in enumerate([col.nombre, col.descripcion, tipo, col.regla], 1):
            celda = ws.cell(fila, i, valor)
            celda.alignment = envolver
            celda.border = BORDE
            if col.formula:
                celda.fill = RELLENO_CALCULADA
    fila += 1
    ws.cell(fila, 1, "* obligatoria").font = Font(italic=True, size=9)


def hoja_conversion(wb: Workbook) -> None:
    ws = wb.create_sheet(HOJA_CONVERSION)
    ws["A1"] = "Conversión de escala literal a numérica"
    ws["A1"].font = ESTILO_TITULO
    ws["A2"] = (
        "PROVISIONAL (decisión D1): se confirma con la IE y se justifica en la tesis. "
        "Escala literal de EBR según RVM N.° 094-2020-MINEDU."
    )
    ws["A2"].font = Font(italic=True, color="C00000")
    for i, cab in enumerate(["Literal", "Valor", "Descripción"], start=1):
        celda = ws.cell(4, i, cab)
        celda.font = ESTILO_ENCABEZADO
        celda.fill = RELLENO_ENCABEZADO
    for fila, (literal, valor, desc) in enumerate(CONVERSION_LITERAL, start=5):
        ws.cell(fila, 1, literal)
        ws.cell(fila, 2, valor)
        ws.cell(fila, 3, desc)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["C"].width = 20
    proteger(ws)


def construir(ficha: Ficha, filas: int) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    escribir_metadatos(ws, ficha)
    escribir_tabla(ws, ficha, filas)
    proteger(ws)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{FILA_ENCABEZADO}:{FILA_ENCABEZADO}"

    hoja_instrucciones(wb, ficha)
    if ficha.con_escala:
        hoja_conversion(wb)
    wb.active = wb.index(ws)
    return wb


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--filas", type=int, default=150, help="Filas de captura por ficha"
    )
    parser.add_argument(
        "--out",
        type=str,
        default=os.path.dirname(os.path.abspath(__file__)),
        help="Carpeta de salida (por defecto, la del script)",
    )
    args = parser.parse_args()

    os.makedirs(args.out, exist_ok=True)
    for ficha in fichas():
        ruta = os.path.join(args.out, ficha.archivo)
        construir(ficha, args.filas).save(ruta)
        print(f"[OK] {ruta} ({len(ficha.columnas)} columnas, {args.filas} filas)")
    print(
        "\nPlantillas vacías. Las fichas llenas se guardan en data/raw/ (no se suben)."
    )


if __name__ == "__main__":
    main()
