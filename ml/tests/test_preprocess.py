"""Pruebas de ml/preprocess.py con fichas de prueba (no son datos reales)."""

import preprocess
import pytest
from conftest import NOTAS_142, PLANTILLAS, alumno
from openpyxl import load_workbook
from preprocess import (
    AREAS,
    INCOMPLETA,
    INCONSISTENTE,
    TRASLADO,
    ErrorValidacion,
    calcular_indicadores,
    consolidar,
)


def asistencia(dp=95, da=85, **kw):
    return alumno(**kw, dias_programados=dp, dias_asistidos=da)


def reuniones(rt=4, ra=1, **kw):
    return alumno(**kw, reuniones_programadas=rt, reuniones_asistidas=ra)


def notas(**kw):
    return alumno(**kw, **NOTAS_142)


def motivo_de(excluidos, codigo):
    fila = excluidos[excluidos["codigo"] == codigo]
    assert len(fila) == 1, f"{codigo} no está excluido"
    return fila.iloc[0]["motivo"], fila.iloc[0]["detalle"]


# --- Fórmulas de la tesis -----------------------------------------------------


def test_formulas_de_la_tesis():
    assert calcular_indicadores(142, 10, 85, 95, 1, 4) == {
        "promedio": 14.2,
        "pct_asistencia": 89.47,
        "pct_reuniones": 25.0,
    }


@pytest.mark.parametrize(
    "argumentos, mensaje",
    [
        ((142, 0, 85, 95, 1, 4), "n_notas"),
        ((142, 10, 85, 0, 1, 4), "dias_programados"),
        ((142, 10, 85, 95, 1, 0), "reuniones_programadas"),
        ((142, 10, 96, 95, 1, 4), "dias_asistidos > dias_programados"),
        ((142, 10, 85, 95, 5, 4), "reuniones_asistidas > reuniones_programadas"),
        ((250, 10, 85, 95, 1, 4), "fuera de la escala"),
    ],
)
def test_reglas_de_validacion(argumentos, mensaje):
    with pytest.raises(ErrorValidacion, match=mensaje):
        calcular_indicadores(*argumentos)


def test_areas_coinciden_con_la_plantilla():
    hoja = load_workbook(PLANTILLAS / "ficha_1_rendimiento.xlsx")["Datos"]
    encabezados = [c.value for c in hoja[preprocess.FILA_ENCABEZADO]]
    assert [a for a in encabezados if a in AREAS] == AREAS
    assert len(AREAS) == 10


# --- Consolidación ------------------------------------------------------------


def test_lote_valido(lote):
    carpeta = lote(
        [notas(), notas(codigo="EST-002", grado=3)],
        [asistencia(), asistencia(codigo="EST-002", grado=3)],
        [reuniones(), reuniones(codigo="EST-002", grado=3)],
    )
    incluidos, excluidos, resumen = consolidar(carpeta)
    assert excluidos.empty
    assert list(incluidos["codigo"]) == ["EST-001", "EST-002"]
    fila = incluidos.iloc[0]
    assert (fila["promedio"], fila["pct_asistencia"], fila["pct_reuniones"]) == (
        14.2,
        89.47,
        25.0,
    )
    assert list(incluidos["grupo"]) == ["experimental", "control"]
    assert "deserto" not in incluidos.columns  # 2026 no lleva deserto
    assert resumen["incluidos_por_grupo"] == {"experimental": 1, "control": 1}


def test_area_exonerada_no_cuenta(lote):
    sin_religion = {k: v for k, v in NOTAS_142.items() if k != "educacion_religiosa"}
    carpeta = lote([alumno(**sin_religion)], [asistencia()], [reuniones()])
    incluidos, _, _ = consolidar(carpeta)
    assert incluidos.iloc[0]["n_notas"] == 9
    assert incluidos.iloc[0]["promedio"] == round(129 / 9, 2)


def test_criterios_de_exclusion(lote):
    carpeta = lote(
        [
            notas(),  # EST-001 válido
            notas(codigo="EST-002"),  # falta en ficha 3
            notas(codigo="EST-003"),  # traslado
            alumno(codigo="EST-004", matematica=25, comunicacion=14),  # nota inválida
            notas(codigo="EST-005"),  # DA > DP
            notas(codigo="Juan Perez"),  # no anonimizado
            notas(codigo="EST-007"),  # grado distinto entre fichas
            notas(codigo="EST-008"),  # repetido en ficha 2
            alumno(codigo="EST-009"),  # sin notas
        ],
        [
            asistencia(),
            asistencia(codigo="EST-002"),
            asistencia(
                codigo="EST-003", excluido="sí", motivo_exclusion="traslado definitivo"
            ),
            asistencia(codigo="EST-004"),
            asistencia(codigo="EST-005", dp=95, da=100),
            asistencia(codigo="Juan Perez"),
            asistencia(codigo="EST-007", grado=3),
            asistencia(codigo="EST-008"),
            asistencia(codigo="EST-008"),
            asistencia(codigo="EST-009"),
        ],
        [
            reuniones(),
            reuniones(codigo="EST-003"),
            reuniones(codigo="EST-004"),
            reuniones(codigo="EST-005"),
            reuniones(codigo="Juan Perez"),
            reuniones(codigo="EST-007"),
            reuniones(codigo="EST-008"),
            reuniones(codigo="EST-009"),
        ],
    )
    incluidos, excluidos, resumen = consolidar(carpeta)

    assert list(incluidos["codigo"]) == ["EST-001"]
    assert motivo_de(excluidos, "EST-002") == (
        INCOMPLETA,
        "falta ficha 3 (apoyo familiar)",
    )
    assert motivo_de(excluidos, "EST-003")[0] == TRASLADO
    assert motivo_de(excluidos, "EST-004")[0] == INCONSISTENTE
    assert "fuera de 0-20" in motivo_de(excluidos, "EST-004")[1]
    assert motivo_de(excluidos, "EST-005") == (
        INCONSISTENTE,
        "dias_asistidos > dias_programados",
    )
    assert "EST-###" in motivo_de(excluidos, "Juan Perez")[1]
    assert "grado distinto" in motivo_de(excluidos, "EST-007")[1]
    assert "repetido en ficha 2" in motivo_de(excluidos, "EST-008")[1]
    assert motivo_de(excluidos, "EST-009") == (INCOMPLETA, "falta notas (ficha 1)")
    assert resumen["excluidos"] == 8
    assert resumen["excluidos_por_motivo"] == {
        TRASLADO: 1,
        INCOMPLETA: 2,
        INCONSISTENTE: 5,
    }


def test_notas_con_decimales(lote):
    carpeta = lote(
        [alumno(matematica=14.5, comunicacion=15.5)], [asistencia()], [reuniones()]
    )
    incluidos, _, _ = consolidar(carpeta)
    assert incluidos.iloc[0]["promedio"] == 15.0


def test_fichas_prellenadas(tmp_path):
    """El juego de --prellenar: 70 códigos, 35 por grado, y los no usados no
    cuentan como excluidos."""
    import subprocess
    import sys

    script = PLANTILLAS / "generar_fichas.py"
    subprocess.run(
        [
            sys.executable,
            str(script),
            "--prellenar",
            "2026_pre",
            "--out",
            str(tmp_path),
        ],
        check=True,
        capture_output=True,
    )
    hoja = load_workbook(tmp_path / "ficha_2_asistencia.xlsx")["Datos"]
    filas = [
        (hoja.cell(r, 1).value, hoja.cell(r, 2).value, hoja.cell(r, 5).value)
        for r in range(preprocess.FILA_INICIO, preprocess.FILA_INICIO + 70)
    ]
    assert filas[0] == ("EST-001", 4, "pre") and filas[34][1] == 4
    assert filas[35] == ("EST-036", 3, "pre") and filas[69][0] == "EST-070"
    assert hoja.cell(preprocess.FILA_INICIO + 70, 1).value is None

    # Se rellenan 2 alumnos completos y a un tercero solo la sección
    from conftest import llenar_prellenada

    llenar_prellenada(tmp_path, 1, {"EST-001": NOTAS_142, "EST-036": NOTAS_142})
    llenar_prellenada(
        tmp_path,
        2,
        {
            "EST-001": {"seccion": "A", "dias_programados": 45, "dias_asistidos": 40},
            "EST-036": {"seccion": "B", "dias_programados": 45, "dias_asistidos": 45},
            "EST-002": {"seccion": "A"},
        },
    )
    llenar_prellenada(
        tmp_path,
        3,
        {
            "EST-001": {"reuniones_programadas": 2, "reuniones_asistidas": 1},
            "EST-036": {"reuniones_programadas": 2, "reuniones_asistidas": 2},
        },
    )
    llenar_prellenada(
        tmp_path, 1, {"EST-001": {"seccion": "A"}, "EST-036": {"seccion": "B"}}
    )
    llenar_prellenada(
        tmp_path, 3, {"EST-001": {"seccion": "A"}, "EST-036": {"seccion": "B"}}
    )

    incluidos, excluidos, resumen = consolidar(tmp_path)
    assert list(incluidos["codigo"]) == ["EST-001", "EST-036"]
    assert list(incluidos["grupo"]) == ["experimental", "control"]
    assert list(excluidos["codigo"]) == ["EST-002"]  # empezado pero incompleto
    assert len(resumen["codigos_sin_usar"]) == 67
    assert resumen["registros"] == 3


def test_historico_requiere_deserto(lote):
    carpeta = lote(
        [notas(anio=2025, deserto=1), notas(codigo="EST-002", anio=2025)],
        [asistencia(anio=2025), asistencia(codigo="EST-002", anio=2025)],
        [reuniones(anio=2025), reuniones(codigo="EST-002", anio=2025)],
    )
    incluidos, excluidos, _ = consolidar(carpeta)
    assert list(incluidos["codigo"]) == ["EST-001"]
    assert incluidos.iloc[0]["deserto"] == 1
    assert motivo_de(excluidos, "EST-002") == (
        INCOMPLETA,
        "falta deserto (obligatorio en el histórico)",
    )


def test_deserto_vacio_en_2026(lote):
    carpeta = lote([notas(deserto=0)], [asistencia()], [reuniones()])
    _, excluidos, _ = consolidar(carpeta)
    assert "vacío en 2026" in motivo_de(excluidos, "EST-001")[1]


def test_anio_desde_el_encabezado(lote):
    sin_anio = [{k: v for k, v in f.items() if k != "anio"} for f in (notas(),)]
    carpeta = lote(
        sin_anio,
        [{k: v for k, v in asistencia().items() if k != "anio"}],
        [{k: v for k, v in reuniones().items() if k != "anio"}],
        anio=2026,
    )
    incluidos, _, _ = consolidar(carpeta)
    assert incluidos.iloc[0]["anio"] == 2026


def test_salida_solo_en_data_processed(tmp_path):
    preprocess.comprobar_salida(preprocess.RAIZ / "data/processed/x.csv")
    preprocess.comprobar_salida(tmp_path / "x.csv")  # fuera del repo: permitido
    with pytest.raises(SystemExit, match="data/processed"):
        preprocess.comprobar_salida(preprocess.RAIZ / "ml/x.csv")
